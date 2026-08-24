from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Software_Maintenance_Handover_Guide_v1.1.0.docx"
NAVY = "163A5F"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5B6573"
WHITE = "FFFFFF"
INK = "17202A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_callout(doc, title, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            for p in cells[index].paragraphs:
                p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.38)
section.footer_distance = Inches(0.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18
for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 9),
    ("Heading 2", 13, BLUE, 13, 6),
    ("Heading 3", 11.5, NAVY, 9, 4),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.18

header = section.header
hp = header.paragraphs[0]
hp.text = "HIRA MODEL SCHOOL | SOFTWARE MAINTENANCE HANDOVER"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.runs[0].font.name = "Calibri"
hp.runs[0].font.size = Pt(8.5)
hp.runs[0].font.bold = True
hp.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
add_page_number(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("MAINTENANCE HANDOVER GUIDE")
r.font.name = "Calibri"
r.font.size = Pt(11)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Hira Model School\nManagement System")
r.font.name = "Calibri"
r.font.size = Pt(29)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(24)
r = p.add_run("Technical, operational, security and release information for the incoming software maintenance provider")
r.font.name = "Calibri"
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(MUTED)

add_matrix(doc, ["Document", "Current baseline"], [
    ("Application version", "1.1.0"),
    ("Database schema", "23"),
    ("Technology", "Python 3.12, PySide6, SQLite, PyInstaller, Inno Setup"),
    ("Validated test baseline", "325 tests passed; 1 intentionally skipped"),
    ("Deployment package", "HiraModelSchoolSetup_1.1.0.exe"),
    ("Owner-only package", "HiraModelSchoolLicenseGenerator_1.1.0_OWNER_ONLY.zip"),
], [2700, 6660])
add_callout(doc, "Purpose", "This document enables a competent software house to understand, support, modify, test and release the system without relying on the original developer for routine maintenance.")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Prepared for controlled maintenance transfer")
r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph("Date: 10 August 2026")
p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

page_break(doc)
doc.add_heading("1. Handover scope and responsibilities", level=1)
doc.add_paragraph("The maintenance provider receives the technical materials required to diagnose defects, implement approved changes, protect customer data and create tested releases. Commercial ownership and licence authority remain subject to the signed maintenance agreement.")
doc.add_heading("1.1 Materials to provide", level=2)
for text in (
    "Private Git repository or complete source package, including app, tests, tools, installer specifications and source manifest.",
    "Requirements file, Python version, PyInstaller specifications and Inno Setup script.",
    "Existing technical reference, software requirements, business rules, installation guide and user operations manual.",
    "A demonstration database containing fake data only, plus a representative encrypted backup for recovery testing.",
    "Current customer installer and one test licence created specifically for acceptance testing.",
    "A written list of supported customers, installed versions and support contacts, without unnecessary student information.",
): add_bullet(doc, text)
doc.add_heading("1.2 Responsibility matrix", level=2)
add_matrix(doc, ["Area", "Owner / school", "Maintenance provider"], [
    ("Business-policy decisions", "Approve", "Recommend; do not assume"),
    ("Bug diagnosis and correction", "Report and verify", "Own implementation"),
    ("Source control and releases", "Receive release record", "Own and preserve history"),
    ("Customer data", "Data controller", "Confidential processor/support party"),
    ("Backups before upgrades", "Confirm availability", "Verify and document"),
    ("Licence generation", "As defined by contract", "Only if expressly authorized"),
], [2100, 3000, 4260])
add_callout(doc, "No implied commercial transfer", "Possession of source code does not by itself transfer copyright, resale rights, customer contracts or authority to issue licences. These rights must be stated in the signed agreement.", LIGHT_GRAY)

page_break(doc)
doc.add_heading("2. Repository and development environment", level=1)
doc.add_heading("2.1 Important project locations", level=2)
add_matrix(doc, ["Location", "Purpose"], [
    ("app/", "Application UI, database initialization, services and business logic"),
    ("tests/", "Unit, integration and regression protection"),
    ("tools/", "Test runners and deployment preflight"),
    ("installer/", "Inno Setup customer-installer definition"),
    ("HiraModelSchool.spec", "Customer PyInstaller build; excludes private licence generator"),
    ("HiraModelSchoolLicenseGenerator.spec", "Owner-only licence generator build"),
    ("SOURCE_PACKAGE_MANIFEST.json", "Source package integrity hashes"),
    ("docs/", "Maintenance notes and regression documentation"),
], [3150, 6210])
doc.add_heading("2.2 Development setup", level=2)
for text in (
    "Install 64-bit Python 3.12 on Windows 11 or a supported Windows release.",
    "Create a virtual environment: python -m venv venv.",
    "Activate it and install dependencies: venv\\Scripts\\python.exe -m pip install -r requirements.txt.",
    "Use QT_QPA_PLATFORM=offscreen when running UI tests without an interactive display.",
    "Run the source application with venv\\Scripts\\python.exe run.py.",
): add_number(doc, text)
doc.add_heading("2.3 Generated files", level=2)
doc.add_paragraph("The build/, dist/, __pycache__/ and .ruff_cache/ folders are generated and may be removed when no build is running. Keep installer_output/ for customer delivery and keep the owner-only ZIP in restricted storage.")

page_break(doc)
doc.add_heading("3. Non-negotiable business rules", level=1)
doc.add_paragraph("Changes affecting the following rules require written business approval before implementation. Technical corrections that preserve these rules may be implemented directly.")
doc.add_heading("3.1 Protected operational rules", level=2)
for title, text in (
    ("Academic session", "The fee session follows April through March. Fees from different sessions remain separately selectable and historically identifiable."),
    ("Dropped students", "A student owes fees through the month in which the drop becomes effective. Later unpaid operational rows must not remain payable."),
    ("Reactivation", "A reactivated student receives a valid current roll according to capacity and roll rules. Historical status and fee gaps remain preserved."),
    ("Archive", "Archiving releases the positive roll number, keeps paid and waived history, removes unpaid operational rows and preserves those removals in the archive audit."),
    ("Graduation", "Graduation and partial graduation depend on the applicable fee state and preserve historical evidence."),
    ("Sibling relationships", "Sibling discounts depend on eligible active relationships and effective event months. Archiving, dropping and graduation must update eligibility without rewriting receipt-backed history."),
    ("Receipts", "Receipts are immutable snapshots. Fee Status History includes previous paid/waived months and the months paid on the current receipt."),
    ("Roll numbers", "Released positive roll numbers may be reused; historical searches must still show the former released number while internal sentinels remain unique."),
):
    doc.add_heading(title, level=3)
    doc.add_paragraph(text)
add_callout(doc, "Approval boundary", "Ask the school before changing fees, discount percentages, academic-session boundaries, receipt content, archive effects, promotion policy, roll allocation or licence policy.")

page_break(doc)
doc.add_heading("4. Database, backups and recovery", level=1)
doc.add_heading("4.1 Database controls", level=2)
for text in (
    "Never edit a customer database manually without first creating and verifying an encrypted backup.",
    "All schema changes must use forward migrations and preserve existing records.",
    "Do not reduce CURRENT_SCHEMA_VERSION or bypass schema-integrity checks.",
    "Keep write operations transactional; a failed operation must not leave partial student, fee or finance changes.",
    "Never package live .db, .db-wal, .db-shm or backup files inside the installer.",
): add_bullet(doc, text)
doc.add_heading("4.2 Backup security", level=2)
doc.add_paragraph("Backups use authenticated encryption and a recovery password. The local Windows key store supports automatic backup operations. The recovery password must be stored separately from customer computers.")
doc.add_heading("4.3 Restore procedure", level=2)
for text in (
    "Confirm the selected backup belongs to the intended school and inspect its date and size.",
    "Authenticate the encrypted backup using the local key or recovery password.",
    "Allow the application to validate and stage the restore; do not replace the live database while the application is running.",
    "Restart so the staged restore completes before ordinary database connections open.",
    "Confirm the pre-restore emergency backup and validate student, fee and finance totals.",
): add_number(doc, text)
add_callout(doc, "Damaged local key store", "Use Backup Management > Recover from Backup, select a valid encrypted .hmsbackup file and enter its recovery password. The damaged key store is retained as evidence and the recovered key is protected for the current Windows account.", LIGHT_GRAY)

page_break(doc)
doc.add_heading("5. Licensing and secrets", level=1)
doc.add_heading("5.1 Package separation", level=2)
add_matrix(doc, ["Package", "Permitted contents", "Distribution"], [
    ("Customer installer", "Application, public verification key, required runtime files", "May be supplied to authorized customers"),
    ("Owner-only generator", "Licence generator and encrypted private signing key", "Restricted to authorized licence administrator"),
], [2250, 4260, 2850])
doc.add_heading("5.2 Prohibited disclosures", level=2)
for text in (
    "Do not send the owner-only ZIP, signing-key password or recovery password to customers.",
    "Do not commit decrypted private keys, customer databases, licences, recovery passwords or credentials to Git.",
    "Do not place private licence-generation modules or encrypted signing-key material in the customer specification.",
    "Transfer licence authority only under a written agreement naming the responsible people and revocation procedure.",
): add_bullet(doc, text)
doc.add_heading("5.3 Licence time integrity", level=2)
doc.add_paragraph("The application records a Windows-protected last trusted time. A material clock rollback produces a warning. Maintenance must preserve signature verification, machine binding, start/expiry validation and the trusted-time warning.")
add_callout(doc, "Recommended governance", "Create a new signing key for the maintenance company if it will issue licences independently. Retain the original owner key privately and document the cutover date.")

page_break(doc)
doc.add_heading("6. Testing, building and releasing", level=1)
doc.add_heading("6.1 Mandatory test gate", level=2)
doc.add_paragraph("Before every release, run:")
add_callout(doc, "Full test command", "venv\\Scripts\\python.exe tools\\run_all_tests.py", LIGHT_GRAY)
doc.add_paragraph("Current expected baseline: 325 tests pass and 1 test is intentionally skipped. Any new failure must be understood and corrected; do not weaken a test merely to obtain a green result.")
doc.add_heading("6.2 Deployment preflight", level=2)
add_callout(doc, "Preflight command", "venv\\Scripts\\python.exe tools\\deployment_preflight.py", LIGHT_GRAY)
doc.add_paragraph("The preflight must report zero warnings/approval items before customer delivery. It checks required files, versions, package separation, runtime-data exclusion, timestamps and Python compilation.")
doc.add_heading("6.3 Build sequence", level=2)
for text in (
    "Run the complete test suite and deployment preflight.",
    "Build the customer application: venv\\Scripts\\python.exe -m PyInstaller --clean --noconfirm HiraModelSchool.spec.",
    "Compile installer\\HiraModelSchoolInstaller.iss with Inno Setup ISCC.exe.",
    "Build the owner tool separately with HiraModelSchoolLicenseGenerator.spec.",
    "Package the owner build into HiraModelSchoolLicenseGenerator_1.1.0_OWNER_ONLY.zip.",
    "Verify hashes, ensure the customer build contains no encrypted_signing_key.json, and record release notes.",
): add_number(doc, text)
doc.add_heading("6.4 Digital signing", level=2)
doc.add_paragraph("Version 1.1.0 executables are currently not Authenticode-signed. A future provider should obtain a trusted code-signing certificate, sign the application and installer, timestamp signatures, and verify them before publication.")

page_break(doc)
doc.add_heading("7. Change and incident management", level=1)
doc.add_heading("7.1 Change workflow", level=2)
for text in (
    "Create a ticket describing the defect or requested behavior, affected customer/version and reproducible evidence.",
    "Classify the change as technical correction or business-rule change.",
    "Obtain written approval for business-rule changes before coding.",
    "Create or update regression tests, implement the smallest safe change and review the database impact.",
    "Run the full test and preflight gates, update documentation and produce versioned release notes.",
    "Back up the customer database, install the update, verify core workflows and document completion.",
): add_number(doc, text)
doc.add_heading("7.2 Severity and response expectations", level=2)
add_matrix(doc, ["Severity", "Examples", "Expected handling"], [
    ("Critical", "Cannot start; data loss; restore/licence failure; incorrect financial totals", "Immediate containment, preserve evidence, restore service safely"),
    ("High", "Payments blocked; wrong fee eligibility; student lifecycle corruption", "Prioritized diagnosis and tested corrective release"),
    ("Normal", "UI defect, report formatting, non-blocking workflow issue", "Schedule into maintenance release"),
    ("Enhancement", "New report, changed business rule, additional workflow", "Estimate and obtain written approval"),
], [1200, 4560, 3600])
doc.add_heading("7.3 Support evidence", level=2)
doc.add_paragraph("Collect the application version, Windows version, exact error, reproduction steps, relevant timestamps and sanitized screenshots. Obtain databases or backups only through an approved secure channel and record who accessed them.")

page_break(doc)
doc.add_heading("8. Security and privacy requirements", level=1)
for text in (
    "Treat student names, CNICs, dates of birth, contact details, addresses and financial records as confidential personal data.",
    "Use fake or anonymized data for development, demonstrations and ordinary testing.",
    "Encrypt customer databases and backups during transfer; do not use ordinary messaging attachments for sensitive data.",
    "Limit repository, owner-generator and production-backup access to named personnel using individual accounts.",
    "Record security-sensitive events and preserve audit evidence rather than overwriting it.",
    "Do not disable password strength, persistent throttling, restore authentication, backup encryption or licence verification to simplify support.",
    "Remove copied customer data from support systems after the agreed retention period and confirm deletion.",
): add_bullet(doc, text)
doc.add_heading("8.1 Release security checklist", level=2)
add_matrix(doc, ["Check", "Required result"], [
    ("Private signing key in customer build", "Absent"),
    ("Live database or backup in package", "Absent"),
    ("Automated tests", "Pass"),
    ("Deployment preflight", "Pass with zero warnings"),
    ("Installer hash recorded", "Yes"),
    ("Database backup verified before upgrade", "Yes"),
    ("Release notes and rollback procedure", "Available"),
], [4050, 5310])

page_break(doc)
doc.add_heading("9. Maintenance-provider acceptance checklist", level=1)
doc.add_paragraph("The incoming provider should demonstrate each item before accepting responsibility. Record the date, evidence and responsible engineer in the formal handover record.")
for text in (
    "Repository received and source manifest verified.",
    "Development environment created from requirements.txt.",
    "Application starts from source using an isolated demonstration database.",
    "All automated tests and deployment preflight pass.",
    "Customer installer is rebuilt from source and installs successfully on a clean Windows test machine.",
    "A test licence is generated separately and activates the customer installation.",
    "Admission, student search, drop, reactivation, archive, promotion and graduation are demonstrated.",
    "Fee generation, multi-session payment, receipt printing and Fee Status History are demonstrated.",
    "Encrypted backup, restore, damaged-key recovery and pre-restore safety backup are demonstrated.",
    "Owner-only signing material is confirmed absent from the customer package.",
    "Support contacts, severity targets, customer list and secure data-transfer procedure are agreed.",
    "Latest source and release artifacts will be returned if the maintenance agreement ends.",
): add_bullet(doc, "☐ " + text)

doc.add_heading("9.1 Formal acceptance", level=2)
add_matrix(doc, ["Field", "Completion"], [
    ("Maintenance provider", "____________________________________________"),
    ("Lead engineer", "____________________________________________"),
    ("Repository/version accepted", "____________________________________________"),
    ("Test and build evidence", "____________________________________________"),
    ("Outstanding exceptions", "____________________________________________"),
    ("Accepted by / signature", "____________________________________________"),
    ("Date", "____________________________________________"),
], [3000, 6360])

page_break(doc)
doc.add_heading("Appendix A. Deliverable inventory", level=1)
add_matrix(doc, ["Deliverable", "Handling classification"], [
    ("Complete source repository", "Confidential - maintenance team"),
    ("Existing requirements, technical and business documents", "Confidential - maintenance team"),
    ("Demonstration database with fake data", "Internal testing"),
    ("Customer installer", "Customer distributable"),
    ("Owner-only licence generator", "Restricted secret material"),
    ("Recovery password", "Restricted; separate secure channel"),
    ("Customer database/backup", "Highly confidential; only when necessary"),
    ("Release hashes and notes", "Maintenance record"),
], [4800, 4560])
doc.add_heading("Appendix B. Existing documents to retain", level=1)
for text in (
    "Full Detailed Technical Reference",
    "Hira Model School Software Requirements Specification",
    "Technical Design Document",
    "Hira Model School Business Rules and Change Record",
    "User Operations Manual",
    "Technician Installation and Activation Guide",
    "docs/REGRESSION_TESTS.md and other current corrective-maintenance reports",
): add_bullet(doc, text)
doc.add_heading("Appendix C. Handover contacts", level=1)
doc.add_paragraph("Complete these fields before transfer:")
add_matrix(doc, ["Role", "Name and secure contact"], [
    ("Software owner / authorized representative", "____________________________________________"),
    ("School operational contact", "____________________________________________"),
    ("Maintenance-provider manager", "____________________________________________"),
    ("Lead maintenance engineer", "____________________________________________"),
    ("Licence administrator", "____________________________________________"),
    ("Emergency data-recovery contact", "____________________________________________"),
], [3900, 5460])

doc.core_properties.title = "Hira Model School Management System - Software Maintenance Handover Guide"
doc.core_properties.subject = "Maintenance provider technical and operational handover"
doc.core_properties.author = "Hira Model School"
doc.core_properties.keywords = "maintenance, handover, deployment, licensing, backups, testing"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
